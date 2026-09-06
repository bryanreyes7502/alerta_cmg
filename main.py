import csv
import glob
import hashlib
import io
import json
import os
import smtplib
import time
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================
# CONFIGURACIÓN
# ============================================================
TIMEZONE = ZoneInfo("America/Santiago")

EMAIL_FROM = "alertas.automaticas@rasopower.cl"
EMAIL_TO = [
    "ingenieria@pmgpinares.cl",
    "bryan.reyes7502@gmail.com",
]
SMTP_HOST = "mail.rasopower.cl"
SMTP_PORT = 465

# La contraseña NO se guarda en el código.
# En GitHub Actions se entrega mediante el Secret EMAIL_APP_PASS.
EMAIL_APP_PASS = os.environ.get("EMAIL_APP_PASS", "").strip()

COORDINADOR_URL = "https://www.coordinador.cl/wp-admin/admin-ajax.php"
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/128.0.0.0 Safari/537.36"
)

BASE_DIR = Path(__file__).resolve().parent
DOWNLOAD_DIR = BASE_DIR / "descargas_csv"
EXCEL_DIR = BASE_DIR / "reportes"
STATE_FILE = BASE_DIR / "state" / "state.json"

DOWNLOAD_DIR.mkdir(parents=True, exist_ok=True)
EXCEL_DIR.mkdir(parents=True, exist_ok=True)
STATE_FILE.parent.mkdir(parents=True, exist_ok=True)

# ============================================================
# CENTRALES Y ZONAS
# ============================================================
CENTRALES_SIN_ZONA = set()

CENTRALES_QUE_SIEMPRE = {
    "CONSTITUCION_DIESEL",
    "MAULE_DIESEL",
    "SANJAVIER-1_DIESEL",
    "SANJAVIER-2_DIESEL",
}

CENTRALES_CONDICIONALES = {
    "CELCO_BL1_COGEN",
    "CELCO_BL1+BL2_COGEN",
    "VINALES_BL1_COGEN",
    "VINALES_BL1+BL2_COGEN",
}

ZONAS_BUSCAR = {
    "LT 66 kV San Javier - Constitución",
    "LT 66 kV Maule - Talca",
}

STATUS_MAP = {
    "APP": "Atraso Proceso de Partida",
    "CI": "Carga Intermedia",
    "EP": "En Pruebas",
    "EXP": "Exportación",
    "FS": "Fuera de Servicio",
    "MT": "Mínimo Técnico",
    "OI": "Operación en Isla",
    "PC": "Plena Carga",
    "PMT": "Proceso a Mínimo Técnico",
    "PP": "Proceso de Partida",
    "PS": "Proceso de Salida",
    "SS": "Seguridad del Sistema",
}

# Mantengo la lógica del script original: CV desactivado.
CV_KEYS = {}
CV_MAP = {}

# ============================================================
# UTILIDADES
# ============================================================
def ahora_local() -> datetime:
    return datetime.now(TIMEZONE)


def log(message: str) -> None:
    print(f"[{ahora_local().strftime('%Y-%m-%d %H:%M:%S %Z')}] {message}", flush=True)


def limpiar_antiguos(carp: Path, patron: str, keep: int = 3) -> None:
    files = sorted(
        carp.glob(patron),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    for file_path in files[keep:]:
        try:
            file_path.unlink()
        except FileNotFoundError:
            pass


def normalizar_estado_env() -> None:
    if not EMAIL_APP_PASS:
        raise RuntimeError(
            "Falta el secreto EMAIL_APP_PASS. Configúralo en "
            "GitHub > Settings > Secrets and variables > Actions."
        )


def cargar_estado() -> set[str]:
    if not STATE_FILE.exists():
        return set()
    try:
        data = json.loads(STATE_FILE.read_text(encoding="utf-8"))
        hashes = data.get("prev_hashes", [])
        return {str(x) for x in hashes}
    except (json.JSONDecodeError, OSError) as exc:
        log(f"No se pudo leer el estado; se inicializará vacío: {exc}")
        return set()


def guardar_estado(hashes: set[str]) -> None:
    payload = {
        "updated_at": ahora_local().isoformat(),
        "prev_hashes": sorted(hashes),
    }
    STATE_FILE.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


# ============================================================
# DESCARGA CSV MEDIANTE CHROME HEADLESS
# ============================================================
COORDINADOR_HOME = "https://www.coordinador.cl/"
COORDINADOR_URL = "https://www.coordinador.cl/wp-admin/admin-ajax.php"

USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/140.0.0.0 Safari/537.36"
)


def crear_driver() -> webdriver.Chrome:
    options = Options()
    options.add_argument("--headless=new")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-gpu")
    options.add_argument("--window-size=1920,1080")
    options.add_argument(f"--user-agent={USER_AGENT}")
    options.add_experimental_option(
        "prefs",
        {
            "download.default_directory": str(DOWNLOAD_DIR),
            "download.prompt_for_download": False,
            "download.directory_upgrade": True,
            "safebrowsing.enabled": True,
        },
    )

    # GitHub-hosted Ubuntu runners ya incluyen Chrome y ChromeDriver.
    # Al indicar explícitamente sus rutas evitamos que Selenium Manager
    # intente descargar otro driver durante la ejecución.
    chrome_binary = os.environ.get("CHROME_BIN", "/usr/bin/google-chrome")
    chromedriver_binary = os.environ.get(
        "CHROMEDRIVER",
        os.environ.get("CHROMEWEBDRIVER", "/usr/bin/chromedriver"),
    )

    if not os.path.exists(chrome_binary):
        candidates = [
            "/usr/bin/google-chrome",
            "/usr/bin/google-chrome-stable",
        ]
        chrome_binary = next((p for p in candidates if os.path.exists(p)), chrome_binary)

    if not os.path.exists(chromedriver_binary):
        candidates = [
            "/usr/local/share/chromedriver-linux64/chromedriver",
            "/usr/bin/chromedriver",
        ]
        chromedriver_binary = next(
            (p for p in candidates if os.path.exists(p)), chromedriver_binary
        )

    options.binary_location = chrome_binary
    return webdriver.Chrome(
        service=Service(chromedriver_binary),
        options=options,
    )


def descargar_csv(fecha: str) -> Path | None:
    params = (
        f"?action=export_energia_csv"
        f"&fecha_inicio={fecha}"
        f"&fecha_termino={fecha}"
        f"&hora_inicio=00:00:00"
        f"&hora_termino=23:59:59"
    )
    url = COORDINADOR_URL + params

    # El endpoint /admin-ajax.php puede bloquear clientes HTTP directos (403).
    # Usamos Chrome headless para reproducir la navegación real del navegador,
    # incluyendo cookies/sesión y ejecución de JavaScript si el sitio lo requiere.
    driver = crear_driver()
    try:
        antes = {p.name for p in DOWNLOAD_DIR.iterdir() if p.is_file()}

        log(f"Abriendo Coordinador para obtener cookies/sesión...")
        driver.get(COORDINADOR_HOME)
        time.sleep(3)

        driver.execute_cdp_cmd(
            "Page.setDownloadBehavior",
            {"behavior": "allow", "downloadPath": str(DOWNLOAD_DIR)},
        )

        log(f"Solicitando CSV al Coordinador para {fecha} mediante Chrome...")
        driver.get(url)

        # Espera a que aparezca el CSV y desaparezca .crdownload.
        deadline = time.time() + 60
        while time.time() < deadline:
            actuales = {p.name for p in DOWNLOAD_DIR.iterdir() if p.is_file()}
            nuevos = actuales - antes

            csvs = [
                DOWNLOAD_DIR / name
                for name in nuevos
                if name.lower().endswith(".csv")
            ]
            parciales = [
                name for name in nuevos
                if name.lower().endswith(".crdownload")
            ]

            if csvs and not parciales:
                src = max(csvs, key=lambda p: p.stat().st_mtime)
                timestamp = ahora_local().strftime("%Y-%m-%d_%H-%M-%S")
                destino = DOWNLOAD_DIR / f"energia_{fecha}_{timestamp}.csv"
                src.rename(destino)
                log(f"CSV descargado: {destino.name} ({destino.stat().st_size:,} bytes)")
                return destino

            time.sleep(1)

        # Diagnóstico útil si el Coordinador respondió una página HTML/403.
        titulo = driver.title
        current_url = driver.current_url
        body = ""
        try:
            body = driver.find_element("tag name", "body").text[:500].replace("\n", " ")
        except Exception:
            pass
        raise RuntimeError(
            "No se encontró el CSV tras 60 s. "
            f"title={titulo!r}, url={current_url!r}, body={body!r}"
        )
    finally:
        driver.quit()


# ============================================================
# EXTRACCIÓN DE HALLAZGOS
# ============================================================
def _leer_csv_rows(csv_path: Path):
    raw = csv_path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = raw.decode(encoding)
            return csv.reader(io.StringIO(text), delimiter=";")
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("csv", raw, 0, 1, "No se pudo decodificar el CSV")


def extraer_hallazgos(csv_path: Path) -> list[dict[str, str]]:
    hall: list[dict[str, str]] = []

    for row in _leer_csv_rows(csv_path):
        if len(row) < 14:
            continue

        cen = row[2].strip('"')
        zona = row[13].strip('"')
        key = cen.upper()

        if key in CENTRALES_SIN_ZONA:
            ok = True
        else:
            ok = (
                key in CENTRALES_QUE_SIEMPRE
                or zona in ZONAS_BUSCAR
                or (
                    key in CENTRALES_CONDICIONALES
                    and any(z in row[12] for z in ZONAS_BUSCAR)
                )
            )

        if not ok:
            continue

        raw_h = row[1].strip('"')
        try:
            h = datetime.strptime(raw_h, "%H:%M:%S").strftime("%H:%M")
        except ValueError:
            h = raw_h[:5]

        sig = row[9].strip('"')
        estado = STATUS_MAP.get(sig, sig)
        carga = row[6].strip('"')
        com = row[12].strip('"')

        hall.append(
            {
                "HORA": h,
                "CENTRAL": cen,
                "CARGA [MW]": carga,
                "ESTADO": estado,
                "COMENTARIO": com,
                "CV [kWh]": "",
                "ZONA DESACOPLE": zona,
            }
        )

    hall.sort(key=lambda x: x["HORA"], reverse=True)
    return hall


# ============================================================
# HASH / NOVEDADES
# ============================================================
def hall_hashes(hall: list[dict[str, str]]) -> set[str]:
    return {
        hashlib.md5(str(sorted(item.items())).encode("utf-8")).hexdigest()
        for item in hall
    }


def hay_novedades(nuevos: set[str], previos: set[str]) -> bool:
    return not nuevos.issubset(previos)


# ============================================================
# EXCEL
# ============================================================
def generar_excel(hall: list[dict[str, str]]) -> Path:
    df = pd.DataFrame(hall)
    ts = ahora_local().strftime("%Y-%m-%d_%H-%M-%S")
    path = EXCEL_DIR / f"reporte_{ts}.xlsx"
    df.to_excel(path, index=False)

    wb = load_workbook(path)
    ws = wb.active
    thin = Side(border_style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    align = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            cell.border = border
            cell.alignment = align
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = max_len + 2

    wb.save(path)
    return path


# ============================================================
# CORREO
# ============================================================
def construir_mensaje(fila: dict[str, str]) -> str:
    return (
        "<p><b>NOVEDADES:</b></p>"
        f"<p><b>HORA:</b> {fila['HORA']} hrs<br>"
        f"<b>CENTRAL:</b> {fila['CENTRAL']}<br>"
        f"<b>ESTADO:</b> {fila['ESTADO']}<br>"
        f"<b>ZONA DESACOPLE:</b> {fila['ZONA DESACOPLE']}</p>"
        "<p>Atte,<br>Sistema Alertas Automáticas.</p>"
    )


def enviar_correo(path_xlsx: Path, hall: list[dict[str, str]]) -> None:
    if not hall:
        return

    fila = hall[0]
    html = construir_mensaje(fila)

    msg = EmailMessage()
    msg["Subject"] = "ALERTA: OPERACIÓN DE CENTRALES"
    msg["From"] = EMAIL_FROM
    msg["To"] = ",".join(EMAIL_TO)
    msg.set_content("Visualiza en un cliente HTML.")
    msg.add_alternative(html, subtype="html")

    with path_xlsx.open("rb") as f:
        msg.add_attachment(
            f.read(),
            maintype="application",
            subtype=(
                "vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            filename=path_xlsx.name,
        )

    log(f"Enviando alerta por correo: {fila['CENTRAL']} / {fila['HORA']}...")
    with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=60) as smtp:
        smtp.login(EMAIL_FROM, EMAIL_APP_PASS)
        smtp.send_message(msg)
    log("Correo enviado correctamente.")


# ============================================================
# PROCESO PRINCIPAL (UNA SOLA EJECUCIÓN)
# ============================================================
def main() -> int:
    normalizar_estado_env()

    fecha = ahora_local().strftime("%Y-%m-%d")
    prev_hashes = cargar_estado()

    csv_path = descargar_csv(fecha)
    if not csv_path:
        log("No hubo CSV. Se termina la ejecución sin generar alerta.")
        return 0

    limpiar_antiguos(DOWNLOAD_DIR, "energia_*.csv", keep=3)

    hall = extraer_hallazgos(csv_path)
    log(f"Hallazgos encontrados: {len(hall)}")

    if not hall:
        return 0

    new_hashes = hall_hashes(hall)

    if not prev_hashes or hay_novedades(new_hashes, prev_hashes):
        xlsx = generar_excel(hall)
        limpiar_antiguos(EXCEL_DIR, "reporte_*.xlsx", keep=3)

        try:
            enviar_correo(xlsx, hall)
        except Exception:
            # No avanzamos el estado si el correo falla. Así la alerta
            # puede reintentarse en la siguiente ejecución.
            raise

        guardar_estado(new_hashes)
        log("Estado actualizado: las novedades actuales ya fueron notificadas.")
    else:
        log("Sin novedades respecto de la última alerta enviada.")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())

